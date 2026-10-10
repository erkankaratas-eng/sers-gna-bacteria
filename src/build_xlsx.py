from common import *
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
T={}
for f in ['s1','s2','s3','s4','s5']: T.update(pd.read_pickle(os.path.join(INTER,f'{f}.pkl')))
dc=pd.DataFrame([[sp,RAWLABEL[s],s,int((strain==s).sum()),f'{X[strain==s].max(1).mean():.0f}'] for sp in SPP for s in SB[sp]],
   columns=['Species','Label in Excel file','Label in manuscript','Spectra','Max intensity in file (min-max scale)'])
README=[
 ('Re-analysis of SERS spectra – Can Gulmus et al.',True),
 ('Input: data/sers_spectra_preprocessed.csv (117 spectra, 9 strains × 13 spectra; baseline-corrected, Savitzky–Golay smoothed and min–max normalized in OPUS 5.5).',False),
 ('Data handling',True),
 ('1. S. aureus spectra were scaled 0–1, P. aeruginosa and E. coli 0–2 in the supplied files. All spectra were re-normalized per spectrum to 0–1 ("minmax"); SNV was used as an alternative preprocessing to test robustness.',False),
 ('2. The first and last 7 data points (constant Savitzky–Golay edge plateau) were removed: 1201 points, 258–1988 cm-1.',False),
 ('3. Strain labels used in the original export files are listed in the Data_check sheet.',False),
 ('Methods',True),
 ('PCA/HCA: mean-centred PCA per species; HCA with Ward linkage on Euclidean distances. HCA stability: 1000 subsamples (10 of 13 spectra per strain), adjusted Rand index (ARI) of the 3-cluster solution vs. strain labels.',False),
 ('Within-strain trend: Spearman correlation between acquisition order and within-strain PC1–PC3 scores.',False),
 ('PERMANOVA (Euclidean): species effect tested with strains as experimental units (exact test over all 280 ways of grouping the 9 strains into three groups of three); strain-within-species effect tested by permuting spectra within species (999 permutations).',False),
 ('Strain classification (within species): PCA (3 PCs, fitted inside every training fold) + LDA or linear SVM (C = 1). Validation: fixed 11/2 split as in manuscript; blocked CV by spectrum ID (blocks 1–4, 5–9, 10–13); repeated 4-fold CV (50×) with label-permutation test (199 permutations). Training-set (resubstitution) results are not reported.',False),
 ('SIMCA: data-driven SIMCA (DD-SIMCA; Pomerantsev 2008, J. Chemometrics 22:601) with 3 PCs per class model, alpha = 0.05, degrees of freedom estimated by method of moments. Class models built from the 11 training spectra of each strain (fixed test spectra as in manuscript). Robustness: 100 random 11/2 splits, 2–4 PCs.',False),
 ('Cross-species challenge: a spectrum is rejected when none of the three strain-specific models of the challenged species accepts it.',False),
 ('Species-level generalization: leave-strain-out over all 27 combinations (one strain per species withheld), PCA (5 PCs) fitted on training strains only, PCA-LDA and linear SVM; exact null distribution from all 280 possible groupings of the 9 strains. Species-level DD-SIMCA built from two strains and tested on the withheld third strain and on the other species.',False),
 ('Notes on uncertainty',True),
 ('95% CIs are Wilson intervals and treat spectra as independent; since all spectra of a strain come from one cultivation and one deposition region, true uncertainty is larger. In leave-strain-out tables each spectrum is predicted 9 times, so no CI is given.',False),
 ('Software: Python 3.12, scikit-learn 1.8, SciPy 1.17 (see requirements.txt). Run: python run_all.py',False),
]
_s7=pd.read_pickle(os.path.join(INTER,'s7.pkl')); _w=pd.read_pickle(os.path.join(INTER,'s8.pkl')); _p=pd.read_pickle(os.path.join(INTER,'s9.pkl')); _bt=pd.read_pickle(os.path.join(INTER,'s10.pkl'))['peaks']; _ne=pd.read_pickle(os.path.join(INTER,'s11.pkl')); _cf=pd.read_pickle(os.path.join(INTER,'s12.pkl')); _cu=pd.read_pickle(os.path.join(INTER,'s15.pkl'))
_g=_s7['gram']; gram=pd.DataFrame([[f'{100*_g[0]:.0f}%',f'{100*_g[1]:.0f}%',f'{100*_g[2]:.0f}%',f'{_g[3]:.2f}']],columns=['Accuracy','Balanced accuracy','Null mean (84 strain selections)','p (exact)'])
_c=_s7['cm']['LSO']; lsocm=pd.DataFrame(_c,index=SPP,columns=SPP).reset_index().rename(columns={'index':'True \\ Predicted'})
win=_w.rename(columns={'centre':'Window centre (cm-1)','from':'From','to':'To','R2_species':'Species (fraction SS)','R2_strain_within':'Strain within species (fraction SS)','strain_acc_blocked':'Strain accuracy, acquisition-order-blocked CV','species_LSO_mean':'Species of withheld strains, mean','species_LSO_best':'Species of withheld strains, best combination'}).round(3)
SHEETS=[('Data_check',dc),('PCA_variance',T['PCA_variance']),('PCA_loading_peaks',T['PCA_loading_peaks']),('HCA',T['HCA']),('Within_strain_trend',T['Drift']),
 ('PERMANOVA',T['PERMANOVA']),('Strain_classification',T['Strain_classification']),('SIMCA_strain_models',T['SIMCA_strain_models']),
 ('SIMCA_test_membership',T['SIMCA_test_membership']),('SIMCA_robustness',T['SIMCA_random_splits']),('Cross_species_challenge',T['Cross_species_challenge']),
 ('Species_leave_strain_out',T['Species_LSO']),('Species_SIMCA_LSO',T['Species_SIMCA_LSO']),('Gram_status_LOSO',gram),('Species_LSO_confusion',lsocm),('Acq_order_correction',_p),('Spectral_windows',win),('Loading_bootstrap',_bt),('Nested_CV',_ne),('Conformal_prediction',_cf),('Accuracy_vs_PCs',_cu['pcs'].round(3)),('Accuracy_vs_spectra',_cu['lc'].round(3))]
wb=Workbook(); ws=wb.active; ws.title='README'
F='Arial'; thin=Side(style='thin',color='999999')
for i,(t,b) in enumerate(README,1):
    c=ws.cell(i,1,t); c.font=Font(name=F,bold=b,size=12 if i==1 else 10); c.alignment=Alignment(wrap_text=True,vertical='top')
ws.column_dimensions['A'].width=130
def put(ws,df,r0=1):
    for j,col in enumerate(df.columns,1):
        c=ws.cell(r0,j,col); c.font=Font(name=F,bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='1F3864'); c.alignment=Alignment(wrap_text=True,vertical='center'); c.border=Border(top=thin,bottom=thin,left=thin,right=thin)
    for i,row in enumerate(df.itertuples(index=False),r0+1):
        for j,v in enumerate(row,1):
            v=v.item() if hasattr(v,'item') else v
            c=ws.cell(i,j,v); c.font=Font(name=F,italic=(df.columns[j-1]=='Species' or df.columns[j-1].endswith('model') and isinstance(v,str) and v in SPP)); c.border=Border(top=thin,bottom=thin,left=thin,right=thin); c.alignment=Alignment(wrap_text=True,vertical='top')
    return r0+len(df)+2
for name,df in SHEETS:
    w=wb.create_sheet(name); put(w,df)
    for j,col in enumerate(df.columns,1):
        L=max([len(str(col))]+[len(str(v)) for v in df[col]]); w.column_dimensions[get_column_letter(j)].width=min(max(12,L*0.9+2),48)
    w.row_dimensions[1].height=45; w.freeze_panes='A2'
w=wb.create_sheet('Strain_confusion_blocked'); r=1
for (sp,kind),cm in T['Strain_confusion_blocked'].items():
    c=w.cell(r,1,f'{sp} – {kind} (acquisition-order-blocked CV; rows = true, columns = predicted)'); c.font=Font(name=F,bold=True); r+=1
    df=cm.reset_index().rename(columns={'index':'True \\ Predicted'}); r=put(w,df,r)
w.column_dimensions['A'].width=22
wb.save(os.path.join(TAB,'reanalysis_results.xlsx')); print('saved', [s for s in wb.sheetnames])
